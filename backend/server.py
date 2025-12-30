from fastapi import FastAPI, APIRouter, Depends, HTTPException, status, UploadFile, File
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os
import logging
from pathlib import Path
from pydantic import BaseModel, Field, validator
from typing import List, Optional, Dict, Any
import uuid
from datetime import datetime, timezone
import jwt
import hashlib
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from io import BytesIO
from fastapi.responses import StreamingResponse
import re

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

# Security
SECRET_KEY = "your-secret-key-here"  # In production, use environment variable
ALGORITHM = "HS256"
security = HTTPBearer()

# MongoDB connection
mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

# Create the main app
app = FastAPI(title="Sistema de Registro de Bonos de Panamá")
api_router = APIRouter(prefix="/api")

# Pydantic Models
class UserBase(BaseModel):
    username: str
    email: str
    brokerage_name: str
    is_active: bool = True

class UserCreate(UserBase):
    password: str

class User(UserBase):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))

class AdminCreate(BaseModel):
    username: str
    password: str
    email: str

class Admin(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    username: str
    email: str
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))

class SecurityBase(BaseModel):
    isin_code: Optional[str] = None
    latinex_code: Optional[str] = None
    security_description: str
    coupon: str
    issue_date: str
    maturity_date: str
    status: str = "En Circulación"  # "En Circulación" | "Vencido"
    
    @validator('isin_code', 'latinex_code')
    def at_least_one_code(cls, v, values):
        if not v and not values.get('isin_code') and not values.get('latinex_code'):
            raise ValueError('Al menos un código ISIN o Latinex es requerido')
        return v

class SecurityCreate(SecurityBase):
    pass

class Security(SecurityBase):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))

class HoldingBase(BaseModel):
    filing_date: str
    isin_or_latinex_code: str
    holder_name: str
    holder_id: str
    legal_representative: Optional[str] = None
    amount_held: float
    address: str
    phone: str
    email: str

class HoldingCreate(HoldingBase):
    pass

class Holding(HoldingBase):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    user_id: str
    security_info: Optional[Dict[str, Any]] = None
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))

class LoginRequest(BaseModel):
    username: str
    password: str

class Token(BaseModel):
    access_token: str
    token_type: str
    user_type: str
    user_info: Dict[str, Any]

# Helper functions
def verify_password(plain_password, hashed_password):
    return hashlib.sha256(plain_password.encode()).hexdigest() == hashed_password

def get_password_hash(password):
    return hashlib.sha256(password.encode()).hexdigest()

def create_access_token(data: dict):
    to_encode = data.copy()
    token = jwt.encode(to_encode, SECRET_KEY, algorithm=ALGORITHM)
    return token

def verify_token(credentials: HTTPAuthorizationCredentials = Depends(security)):
    try:
        payload = jwt.decode(credentials.credentials, SECRET_KEY, algorithms=[ALGORITHM])
        return payload
    except jwt.PyJWTError:
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Token inválido",
            headers={"WWW-Authenticate": "Bearer"},
        )

def prepare_for_mongo(data):
    """Convert datetime objects to ISO strings for MongoDB storage"""
    if isinstance(data, dict):
        for key, value in data.items():
            if isinstance(value, datetime):
                data[key] = value.isoformat()
    return data

# Authentication endpoints
@api_router.post("/auth/login", response_model=Token)
async def login(login_data: LoginRequest):
    # Check admin first
    admin = await db.admins.find_one({"username": login_data.username})
    if admin and verify_password(login_data.password, admin["hashed_password"]):
        token_data = {"sub": admin["username"], "user_type": "admin", "user_id": admin["id"]}
        token = create_access_token(token_data)
        return {
            "access_token": token,
            "token_type": "bearer",
            "user_type": "admin",
            "user_info": {"username": admin["username"], "email": admin["email"]}
        }
    
    # Check user
    user = await db.users.find_one({"username": login_data.username})
    if user and verify_password(login_data.password, user["hashed_password"]):
        if not user.get("is_active", True):
            raise HTTPException(status_code=400, detail="Usuario inactivo")
        token_data = {"sub": user["username"], "user_type": "user", "user_id": user["id"]}
        token = create_access_token(token_data)
        return {
            "access_token": token,
            "token_type": "bearer",
            "user_type": "user",
            "user_info": {
                "username": user["username"], 
                "email": user["email"],
                "brokerage_name": user["brokerage_name"]
            }
        }
    
    raise HTTPException(status_code=400, detail="Credenciales incorrectas")

# Admin endpoints
@api_router.post("/admin/create-admin", response_model=Admin)
async def create_admin(admin_data: AdminCreate, token_payload: dict = Depends(verify_token)):
    if token_payload.get("user_type") != "admin":
        raise HTTPException(status_code=403, detail="Acceso denegado")
    
    # Check if admin exists
    existing = await db.admins.find_one({"username": admin_data.username})
    if existing:
        raise HTTPException(status_code=400, detail="El administrador ya existe")
    
    hashed_password = get_password_hash(admin_data.password)
    admin_dict = admin_data.dict()
    del admin_dict["password"]
    admin_obj = Admin(**admin_dict)
    admin_dict = admin_obj.dict()
    admin_dict["hashed_password"] = hashed_password
    admin_dict = prepare_for_mongo(admin_dict)
    
    await db.admins.insert_one(admin_dict)
    return admin_obj

@api_router.get("/admin/admins", response_model=List[Admin])
async def get_admins(token_payload: dict = Depends(verify_token)):
    if token_payload.get("user_type") != "admin":
        raise HTTPException(status_code=403, detail="Acceso denegado")
    
    admins = await db.admins.find().to_list(1000)
    return [Admin(**admin) for admin in admins]

@api_router.put("/admin/admins/{admin_id}", response_model=Admin)
async def update_admin(admin_id: str, admin_data: dict, token_payload: dict = Depends(verify_token)):
    if token_payload.get("user_type") != "admin":
        raise HTTPException(status_code=403, detail="Acceso denegado")
    
    # Check if admin exists
    existing_admin = await db.admins.find_one({"id": admin_id})
    if not existing_admin:
        raise HTTPException(status_code=404, detail="Administrador no encontrado")
    
    # Validate required fields
    if not admin_data.get("email") or not admin_data.get("username"):
        raise HTTPException(status_code=400, detail="Username y email son requeridos")
    
    # Check for duplicate username (excluding current admin)
    if "username" in admin_data:
        duplicate = await db.admins.find_one({"username": admin_data["username"], "id": {"$ne": admin_id}})
        if duplicate:
            raise HTTPException(status_code=400, detail="Ya existe otro administrador con ese nombre de usuario")
    
    # Prepare update data
    update_data = {
        "username": admin_data["username"],
        "email": admin_data["email"],
        "updated_at": datetime.now(timezone.utc).isoformat()
    }
    
    # Update password if provided
    if admin_data.get("password"):
        if len(admin_data["password"]) < 6:
            raise HTTPException(status_code=400, detail="La contraseña debe tener al menos 6 caracteres")
        update_data["hashed_password"] = get_password_hash(admin_data["password"])
    
    # Update admin
    await db.admins.update_one({"id": admin_id}, {"$set": update_data})
    
    # Return updated admin
    updated_admin = await db.admins.find_one({"id": admin_id})
    return Admin(**updated_admin)

@api_router.put("/admin/admins/{admin_id}/password")
async def update_admin_password(admin_id: str, password_data: dict, token_payload: dict = Depends(verify_token)):
    if token_payload.get("user_type") != "admin":
        raise HTTPException(status_code=403, detail="Acceso denegado")
    
    # Validate password data
    new_password = password_data.get("new_password")
    if not new_password or len(new_password) < 6:
        raise HTTPException(status_code=400, detail="La contraseña debe tener al menos 6 caracteres")
    
    # Check if admin exists
    admin = await db.admins.find_one({"id": admin_id})
    if not admin:
        raise HTTPException(status_code=404, detail="Administrador no encontrado")
    
    # Update password
    hashed_password = get_password_hash(new_password)
    await db.admins.update_one(
        {"id": admin_id}, 
        {"$set": {
            "hashed_password": hashed_password,
            "updated_at": datetime.now(timezone.utc).isoformat()
        }}
    )
    
    return {"message": "Contraseña de administrador actualizada exitosamente"}

@api_router.post("/admin/users", response_model=User)
async def create_user(user_data: UserCreate, token_payload: dict = Depends(verify_token)):
    if token_payload.get("user_type") != "admin":
        raise HTTPException(status_code=403, detail="Acceso denegado")
    
    # Check if user exists
    existing = await db.users.find_one({"username": user_data.username})
    if existing:
        raise HTTPException(status_code=400, detail="El usuario ya existe")
    
    hashed_password = get_password_hash(user_data.password)
    user_dict = user_data.dict()
    del user_dict["password"]
    user_obj = User(**user_dict)
    user_dict = user_obj.dict()
    user_dict["hashed_password"] = hashed_password
    user_dict = prepare_for_mongo(user_dict)
    
    await db.users.insert_one(user_dict)
    return user_obj

@api_router.get("/admin/users", response_model=List[User])
async def get_users(token_payload: dict = Depends(verify_token)):
    if token_payload.get("user_type") != "admin":
        raise HTTPException(status_code=403, detail="Acceso denegado")
    
    users = await db.users.find().to_list(1000)
    return [User(**user) for user in users]

@api_router.put("/admin/users/{user_id}", response_model=User)
async def update_user(user_id: str, user_data: UserCreate, token_payload: dict = Depends(verify_token)):
    if token_payload.get("user_type") != "admin":
        raise HTTPException(status_code=403, detail="Acceso denegado")
    
    # Check if user exists
    existing_user = await db.users.find_one({"id": user_id})
    if not existing_user:
        raise HTTPException(status_code=404, detail="Usuario no encontrado")
    
    # Check for duplicate username (excluding current user)
    duplicate = await db.users.find_one({"username": user_data.username, "id": {"$ne": user_id}})
    if duplicate:
        raise HTTPException(status_code=400, detail="Ya existe otro usuario con ese nombre de usuario")
    
    # Prepare update data
    update_data = {
        "username": user_data.username,
        "email": user_data.email,
        "brokerage_name": user_data.brokerage_name,
        "is_active": user_data.is_active,
        "updated_at": datetime.now(timezone.utc).isoformat()
    }
    
    # Update password if provided
    if user_data.password:
        update_data["hashed_password"] = get_password_hash(user_data.password)
    
    # Update user
    await db.users.update_one({"id": user_id}, {"$set": update_data})
    
    # Return updated user
    updated_user = await db.users.find_one({"id": user_id})
    return User(**updated_user)

@api_router.put("/admin/users/{user_id}/password")
async def update_user_password(user_id: str, password_data: dict, token_payload: dict = Depends(verify_token)):
    if token_payload.get("user_type") != "admin":
        raise HTTPException(status_code=403, detail="Acceso denegado")
    
    # Validate password data
    new_password = password_data.get("new_password")
    if not new_password or len(new_password) < 6:
        raise HTTPException(status_code=400, detail="La contraseña debe tener al menos 6 caracteres")
    
    # Check if user exists
    user = await db.users.find_one({"id": user_id})
    if not user:
        raise HTTPException(status_code=404, detail="Usuario no encontrado")
    
    # Update password
    hashed_password = get_password_hash(new_password)
    await db.users.update_one(
        {"id": user_id}, 
        {"$set": {
            "hashed_password": hashed_password,
            "updated_at": datetime.now(timezone.utc).isoformat()
        }}
    )
    
    return {"message": "Contraseña actualizada exitosamente"}

@api_router.delete("/admin/users/{user_id}")
async def delete_user(user_id: str, token_payload: dict = Depends(verify_token)):
    if token_payload.get("user_type") != "admin":
        raise HTTPException(status_code=403, detail="Acceso denegado")
    
    # Check if user exists
    user = await db.users.find_one({"id": user_id})
    if not user:
        raise HTTPException(status_code=404, detail="Usuario no encontrado")
    
    # Check if user has holdings
    holdings_count = await db.holdings.count_documents({"user_id": user_id})
    if holdings_count > 0:
        raise HTTPException(
            status_code=400, 
            detail=f"No se puede eliminar el usuario porque tiene {holdings_count} tenencias registradas"
        )
    
    # Delete user
    await db.users.delete_one({"id": user_id})
    return {"message": "Usuario eliminado exitosamente"}

@api_router.put("/admin/users/{user_id}/toggle")
async def toggle_user_status(user_id: str, token_payload: dict = Depends(verify_token)):
    if token_payload.get("user_type") != "admin":
        raise HTTPException(status_code=403, detail="Acceso denegado")
    
    user = await db.users.find_one({"id": user_id})
    if not user:
        raise HTTPException(status_code=404, detail="Usuario no encontrado")
    
    new_status = not user.get("is_active", True)
    await db.users.update_one({"id": user_id}, {"$set": {"is_active": new_status}})
    return {"message": f"Usuario {'activado' if new_status else 'desactivado'}"}

# Securities (ISIN/Latinex) endpoints
@api_router.post("/admin/securities", response_model=Security)
async def create_security(security_data: SecurityCreate, token_payload: dict = Depends(verify_token)):
    if token_payload.get("user_type") != "admin":
        raise HTTPException(status_code=403, detail="Acceso denegado")
    
    # Check for duplicates
    query = {}
    if security_data.isin_code:
        query["isin_code"] = security_data.isin_code
    if security_data.latinex_code:
        if query:
            query = {"$or": [query, {"latinex_code": security_data.latinex_code}]}
        else:
            query["latinex_code"] = security_data.latinex_code
    
    existing = await db.securities.find_one(query)
    if existing:
        raise HTTPException(status_code=400, detail="Ya existe un valor con ese código")
    
    security_obj = Security(**security_data.dict())
    security_dict = prepare_for_mongo(security_obj.dict())
    await db.securities.insert_one(security_dict)
    return security_obj

@api_router.get("/admin/securities", response_model=List[Security])
async def get_securities(token_payload: dict = Depends(verify_token)):
    if token_payload.get("user_type") != "admin":
        raise HTTPException(status_code=403, detail="Acceso denegado")
    
    securities = await db.securities.find().to_list(1000)
    return [Security(**security) for security in securities]

@api_router.put("/admin/securities/{security_id}", response_model=Security)
async def update_security(security_id: str, security_data: SecurityCreate, token_payload: dict = Depends(verify_token)):
    if token_payload.get("user_type") != "admin":
        raise HTTPException(status_code=403, detail="Acceso denegado")
    
    # Check if security exists
    existing = await db.securities.find_one({"id": security_id})
    if not existing:
        raise HTTPException(status_code=404, detail="Valor no encontrado")
    
    # Check for duplicates (excluding current security)
    query = {"id": {"$ne": security_id}}
    if security_data.isin_code:
        query["$or"] = [{"isin_code": security_data.isin_code}]
    if security_data.latinex_code:
        if "$or" in query:
            query["$or"].append({"latinex_code": security_data.latinex_code})
        else:
            query["$or"] = [{"latinex_code": security_data.latinex_code}]
    
    if "$or" in query:
        duplicate = await db.securities.find_one(query)
        if duplicate:
            raise HTTPException(status_code=400, detail="Ya existe otro valor con ese código ISIN o Latinex")
    
    # Update the security
    update_data = security_data.dict()
    update_data["updated_at"] = datetime.now(timezone.utc).isoformat()
    
    await db.securities.update_one(
        {"id": security_id}, 
        {"$set": update_data}
    )
    
    # Return updated security
    updated_security = await db.securities.find_one({"id": security_id})
    return Security(**updated_security)

@api_router.delete("/admin/securities/{security_id}")
async def delete_security(security_id: str, token_payload: dict = Depends(verify_token)):
    if token_payload.get("user_type") != "admin":
        raise HTTPException(status_code=403, detail="Acceso denegado")
    
    # Check if security exists
    existing = await db.securities.find_one({"id": security_id})
    if not existing:
        raise HTTPException(status_code=404, detail="Valor no encontrado")
    
    # Delete the security
    await db.securities.delete_one({"id": security_id})
    
    return {"message": "Valor eliminado exitosamente"}

@api_router.delete("/admin/securities/clear-all")
async def clear_all_securities(token_payload: dict = Depends(verify_token)):
    if token_payload.get("user_type") != "admin":
        raise HTTPException(status_code=403, detail="Acceso denegado")
    
    # Get count before deletion
    count_before = await db.securities.count_documents({})
    
    # Delete all securities
    result = await db.securities.delete_many({})
    
    return {
        "message": f"Eliminados {result.deleted_count} valores registrados",
        "deleted_count": result.deleted_count,
        "total_before": count_before
    }

@api_router.get("/securities/search/{code}")
async def search_security(code: str, token_payload: dict = Depends(verify_token)):
    """Search security by ISIN or Latinex code"""
    query = {"$or": [{"isin_code": code}, {"latinex_code": code}]}
    security = await db.securities.find_one(query)
    if not security:
        raise HTTPException(status_code=404, detail="Valor no encontrado")
    return Security(**security)

# Holdings endpoints
@api_router.post("/holdings", response_model=Holding)
async def create_holding(holding_data: HoldingCreate, token_payload: dict = Depends(verify_token)):
    if token_payload.get("user_type") != "user":
        raise HTTPException(status_code=403, detail="Acceso denegado")
    
    # Get security info
    query = {"$or": [
        {"isin_code": holding_data.isin_or_latinex_code}, 
        {"latinex_code": holding_data.isin_or_latinex_code}
    ]}
    security = await db.securities.find_one(query)
    
    holding_dict = holding_data.dict()
    holding_dict["user_id"] = token_payload["user_id"]
    if security:
        holding_dict["security_info"] = {
            "isin_code": security.get("isin_code"),
            "latinex_code": security.get("latinex_code"),
            "security_description": security["security_description"],
            "coupon": security["coupon"],
            "maturity_date": security["maturity_date"]
        }
    
    holding_obj = Holding(**holding_dict)
    holding_dict = prepare_for_mongo(holding_obj.dict())
    await db.holdings.insert_one(holding_dict)
    
    # Update master holdings
    await update_master_holdings()
    
    return holding_obj

@api_router.get("/holdings", response_model=List[Holding])
async def get_user_holdings(token_payload: dict = Depends(verify_token)):
    if token_payload.get("user_type") != "user":
        raise HTTPException(status_code=403, detail="Acceso denegado")
    
    holdings = await db.holdings.find({"user_id": token_payload["user_id"]}).to_list(1000)
    return [Holding(**holding) for holding in holdings]

@api_router.get("/admin/holdings", response_model=List[Holding])
async def get_all_holdings(token_payload: dict = Depends(verify_token)):
    if token_payload.get("user_type") != "admin":
        raise HTTPException(status_code=403, detail="Acceso denegado")
    
    holdings = await db.holdings.find().to_list(1000)
    return [Holding(**holding) for holding in holdings]

@api_router.get("/admin/users/{user_id}/holdings", response_model=List[Holding])
async def get_user_holdings_by_admin(user_id: str, token_payload: dict = Depends(verify_token)):
    if token_payload.get("user_type") != "admin":
        raise HTTPException(status_code=403, detail="Acceso denegado")
    
    # Verify user exists
    user = await db.users.find_one({"id": user_id})
    if not user:
        raise HTTPException(status_code=404, detail="Usuario no encontrado")
    
    holdings = await db.holdings.find({"user_id": user_id}).to_list(1000)
    return [Holding(**holding) for holding in holdings]

@api_router.post("/admin/users/{user_id}/holdings", response_model=Holding)
async def create_holding_for_user(user_id: str, holding_data: HoldingCreate, token_payload: dict = Depends(verify_token)):
    if token_payload.get("user_type") != "admin":
        raise HTTPException(status_code=403, detail="Acceso denegado")
    
    # Verify user exists
    user = await db.users.find_one({"id": user_id})
    if not user:
        raise HTTPException(status_code=404, detail="Usuario no encontrado")
    
    # Get security info
    query = {"$or": [
        {"isin_code": holding_data.isin_or_latinex_code}, 
        {"latinex_code": holding_data.isin_or_latinex_code}
    ]}
    security = await db.securities.find_one(query)
    
    holding_dict = holding_data.dict()
    holding_dict["user_id"] = user_id
    if security:
        holding_dict["security_info"] = {
            "isin_code": security.get("isin_code"),
            "latinex_code": security.get("latinex_code"),
            "security_description": security["security_description"],
            "coupon": security["coupon"],
            "maturity_date": security["maturity_date"]
        }
    
    holding_obj = Holding(**holding_dict)
    holding_dict = prepare_for_mongo(holding_obj.dict())
    await db.holdings.insert_one(holding_dict)
    
    # Update master holdings
    await update_master_holdings()
    
    return holding_obj

@api_router.put("/admin/holdings/{holding_id}", response_model=Holding)
async def update_holding_by_admin(holding_id: str, holding_data: HoldingCreate, token_payload: dict = Depends(verify_token)):
    if token_payload.get("user_type") != "admin":
        raise HTTPException(status_code=403, detail="Acceso denegado")
    
    # Check if holding exists
    existing = await db.holdings.find_one({"id": holding_id})
    if not existing:
        raise HTTPException(status_code=404, detail="Tenencia no encontrada")
    
    # Get security info
    query = {"$or": [
        {"isin_code": holding_data.isin_or_latinex_code}, 
        {"latinex_code": holding_data.isin_or_latinex_code}
    ]}
    security = await db.securities.find_one(query)
    
    # Prepare update data
    update_data = holding_data.dict()
    update_data["updated_at"] = datetime.now(timezone.utc).isoformat()
    
    if security:
        update_data["security_info"] = {
            "isin_code": security.get("isin_code"),
            "latinex_code": security.get("latinex_code"),
            "security_description": security["security_description"],
            "coupon": security["coupon"],
            "maturity_date": security["maturity_date"]
        }
    
    # Update holding
    await db.holdings.update_one({"id": holding_id}, {"$set": update_data})
    
    # Update master holdings
    await update_master_holdings()
    
    # Return updated holding
    updated_holding = await db.holdings.find_one({"id": holding_id})
    return Holding(**updated_holding)

@api_router.delete("/admin/holdings/{holding_id}")
async def delete_holding_by_admin(holding_id: str, token_payload: dict = Depends(verify_token)):
    if token_payload.get("user_type") != "admin":
        raise HTTPException(status_code=403, detail="Acceso denegado")
    
    # Check if holding exists
    existing = await db.holdings.find_one({"id": holding_id})
    if not existing:
        raise HTTPException(status_code=404, detail="Tenencia no encontrada")
    
    # Delete holding
    await db.holdings.delete_one({"id": holding_id})
    
    # Update master holdings
    await update_master_holdings()
    
    return {"message": "Tenencia eliminada exitosamente"}

async def update_master_holdings():
    """Update master holdings collection with consolidated data"""
    # Clear existing master holdings
    await db.master_holdings.delete_many({})
    
    # Aggregate holdings by filing_date, security, and holder
    pipeline = [
        {
            "$group": {
                "_id": {
                    "filing_date": "$filing_date",
                    "isin_or_latinex_code": "$isin_or_latinex_code",
                    "holder_name": "$holder_name",
                    "holder_id": "$holder_id"
                },
                "total_amount": {"$sum": "$amount_held"},
                "holdings": {"$push": "$$ROOT"}
            }
        },
        {
            "$sort": {
                "_id.filing_date": 1,
                "_id.isin_or_latinex_code": 1,
                "_id.holder_name": 1
            }
        }
    ]
    
    aggregated = await db.holdings.aggregate(pipeline).to_list(1000)
    
    master_holdings = []
    for item in aggregated:
        master_holding = {
            "id": str(uuid.uuid4()),
            "filing_date": item["_id"]["filing_date"],
            "isin_or_latinex_code": item["_id"]["isin_or_latinex_code"],
            "holder_name": item["_id"]["holder_name"],
            "holder_id": item["_id"]["holder_id"],
            "total_amount": item["total_amount"],
            "holdings_count": len(item["holdings"]),
            "source_holdings": item["holdings"],
            "created_at": datetime.now(timezone.utc).isoformat()
        }
        master_holdings.append(master_holding)
    
    if master_holdings:
        await db.master_holdings.insert_many(master_holdings)

@api_router.post("/admin/import/excel")
async def import_from_excel(file: UploadFile = File(...), token_payload: dict = Depends(verify_token)):
    if token_payload.get("user_type") != "admin":
        raise HTTPException(status_code=403, detail="Acceso denegado")
    
    if not file.filename.endswith(('.xlsx', '.xls')):
        raise HTTPException(status_code=400, detail="Formato de archivo no válido. Use archivos Excel (.xlsx, .xls)")
    
    try:
        # Read Excel file
        contents = await file.read()
        wb = openpyxl.load_workbook(BytesIO(contents))
        
        # Look for "Valores_ISIN" sheet or first sheet
        sheet_name = "Valores_ISIN" if "Valores_ISIN" in wb.sheetnames else wb.sheetnames[0]
        ws = wb[sheet_name]
        
        imported_count = 0
        errors = []
        
        # Expected columns: Código ISIN, Código Latinex, Descripción del Valor, Cupón, Fecha de Emisión, Fecha de Vencimiento
        for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not any(row):  # Skip empty rows
                continue
            
            try:
                isin_code = str(row[0]).strip() if row[0] else None
                latinex_code = str(row[1]).strip() if row[1] else None
                security_description = str(row[2]).strip() if row[2] else ""
                coupon = str(row[3]).strip() if row[3] else ""
                issue_date = str(row[4]).strip() if row[4] else ""
                maturity_date = str(row[5]).strip() if row[5] else ""
                
                # Validate required fields
                if not security_description:
                    errors.append(f"Fila {row_num}: Descripción del valor es requerida")
                    continue
                
                if not isin_code and not latinex_code:
                    errors.append(f"Fila {row_num}: Se requiere al menos un código ISIN o Latinex")
                    continue
                
                if not coupon or not issue_date or not maturity_date:
                    errors.append(f"Fila {row_num}: Cupón, fecha de emisión y fecha de vencimiento son requeridos")
                    continue
                
                # Check for duplicates in database
                query = {}
                if isin_code:
                    query["isin_code"] = isin_code
                if latinex_code:
                    if query:
                        query = {"$or": [query, {"latinex_code": latinex_code}]}
                    else:
                        query["latinex_code"] = latinex_code
                
                existing = await db.securities.find_one(query)
                if existing:
                    errors.append(f"Fila {row_num}: Ya existe un valor con código ISIN '{isin_code}' o Latinex '{latinex_code}'")
                    continue
                
                # Create security object
                security_data = {
                    "isin_code": isin_code if isin_code else None,
                    "latinex_code": latinex_code if latinex_code else None,
                    "security_description": security_description,
                    "coupon": coupon,
                    "issue_date": issue_date,
                    "maturity_date": maturity_date
                }
                
                security_obj = Security(**security_data)
                security_dict = prepare_for_mongo(security_obj.dict())
                
                await db.securities.insert_one(security_dict)
                imported_count += 1
                
            except Exception as e:
                errors.append(f"Fila {row_num}: Error procesando datos - {str(e)}")
                continue
        
        # Prepare response
        response_data = {
            "message": f"Importación completada: {imported_count} valores importados",
            "imported_count": imported_count,
            "total_errors": len(errors),
            "errors": errors[:10] if errors else []  # Limit to first 10 errors
        }
        
        if errors:
            response_data["message"] += f", {len(errors)} errores encontrados"
        
        return response_data
        
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Error procesando archivo: {str(e)}")

@api_router.get("/admin/export/excel")
async def export_to_excel(token_payload: dict = Depends(verify_token)):
    if token_payload.get("user_type") != "admin":
        raise HTTPException(status_code=403, detail="Acceso denegado")
    
    # Create workbook
    wb = openpyxl.Workbook()
    
    # Remove default sheet
    wb.remove(wb.active)
    
    # Style definitions
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    center_alignment = Alignment(horizontal="center", vertical="center")
    
    # Sheet 1: Securities (ISIN Master)
    ws1 = wb.create_sheet("Valores_ISIN")
    securities_headers = ["Código ISIN", "Código Latinex", "Descripción del Valor", 
                         "Cupón", "Fecha de Emisión", "Fecha de Vencimiento"]
    ws1.append(securities_headers)
    
    for cell in ws1[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_alignment
    
    securities = await db.securities.find().to_list(1000)
    for security in securities:
        ws1.append([
            security.get("isin_code", ""),
            security.get("latinex_code", ""),
            security["security_description"],
            security["coupon"],
            security["issue_date"],
            security["maturity_date"]
        ])
    
    # Sheet 2: Individual Holdings
    ws2 = wb.create_sheet("Tenencias_Individuales")
    holdings_headers = ["Fecha de Presentación", "Código ISIN/Latinex", "Nombre del Tenedor", 
                       "ID del Tenedor", "Representante Legal", "Cantidad Tenida", 
                       "Dirección", "Teléfono", "Email", "Casa de Corretaje"]
    ws2.append(holdings_headers)
    
    for cell in ws2[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_alignment
    
    # Get holdings with user info
    pipeline = [
        {
            "$lookup": {
                "from": "users",
                "localField": "user_id",
                "foreignField": "id",
                "as": "user_info"
            }
        },
        {
            "$sort": {
                "filing_date": 1,
                "isin_or_latinex_code": 1,
                "holder_name": 1
            }
        }
    ]
    
    holdings_with_users = await db.holdings.aggregate(pipeline).to_list(1000)
    for holding in holdings_with_users:
        user_info = holding.get("user_info", [{}])[0]
        ws2.append([
            holding["filing_date"],
            holding["isin_or_latinex_code"],
            holding["holder_name"],
            holding["holder_id"],
            holding.get("legal_representative", ""),
            holding["amount_held"],
            holding["address"],
            holding["phone"],
            holding["email"],
            user_info.get("brokerage_name", "")
        ])
    
    # Sheet 3: Master Consolidated Holdings
    ws3 = wb.create_sheet("Tenencias_Consolidadas")
    master_headers = ["Fecha de Presentación", "Código ISIN/Latinex", "Descripción del Valor",
                     "Nombre del Tenedor", "ID del Tenedor", "Cantidad Total", 
                     "Número de Registros"]
    ws3.append(master_headers)
    
    for cell in ws3[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_alignment
    
    master_holdings = await db.master_holdings.find().sort([
        ("filing_date", 1), ("isin_or_latinex_code", 1), ("holder_name", 1)
    ]).to_list(1000)
    
    for master in master_holdings:
        # Get security description
        security_info = None
        if master.get("source_holdings"):
            security_info = master["source_holdings"][0].get("security_info", {})
        
        ws3.append([
            master["filing_date"],
            master["isin_or_latinex_code"],
            security_info.get("security_description", "") if security_info else "",
            master["holder_name"],
            master["holder_id"],
            master["total_amount"],
            master["holdings_count"]
        ])
    
    # Auto-adjust column widths
    for ws in [ws1, ws2, ws3]:
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
    
    # Save to BytesIO
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    # Create filename with timestamp
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"tenencias_bonos_panama_{timestamp}.xlsx"
    
    return StreamingResponse(
        BytesIO(output.read()),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

# Create default admin on startup
@app.on_event("startup")
async def create_default_admin():
    admin_exists = await db.admins.find_one({"username": "admin"})
    if not admin_exists:
        admin_data = {
            "id": str(uuid.uuid4()),
            "username": "admin",
            "email": "admin@sistema.com",
            "hashed_password": get_password_hash("admin123"),
            "created_at": datetime.now(timezone.utc).isoformat()
        }
        await db.admins.insert_one(admin_data)
        print("Admin por defecto creado: usuario 'admin', contraseña 'admin123'")

# Include router
app.include_router(api_router)

# CORS
app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get('CORS_ORIGINS', '*').split(','),
    allow_methods=["*"],
    allow_headers=["*"],
)

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
