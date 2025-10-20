import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

# Create a test Excel file for importing securities
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Valores_ISIN"

# Style definitions
header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
center_alignment = Alignment(horizontal="center", vertical="center")

# Headers
headers = ["Código ISIN", "Código Latinex", "Descripción del Valor", "Cupón", "Fecha de Emisión", "Fecha de Vencimiento"]
ws.append(headers)

# Style headers
for cell in ws[1]:
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = center_alignment

# Sample data
sample_data = [
    ["US698299BK05", "RPMA033620231A", "República de Panamá - Bono Serie A", "8.500%", "2023-01-15", "2033-01-15"],
    ["US698299CK04", "RPMA033620241B", "República de Panamá - Bono Serie B", "7.750%", "2024-02-20", "2034-02-20"],
    ["", "RPME093750429C", "República de Panamá - Bono Local", "9.375%", "2024-04-01", "2029-04-01"],
    ["US698299DK03", "", "República de Panamá - Bono Internacional", "6.250%", "2024-06-15", "2031-06-15"],
    ["US698299EK02", "RPMA033620251E", "República de Panamá - Bono Verde", "5.875%", "2025-01-10", "2035-01-10"]
]

for row_data in sample_data:
    ws.append(row_data)

# Auto-adjust column widths
for column in ws.columns:
    max_length = 0
    column_letter = column[0].column_letter
    for cell in column:
        try:
            if len(str(cell.value)) > max_length:
                max_length = len(str(cell.value))
        except:
            pass
    adjusted_width = min(max_length + 2, 30)
    ws.column_dimensions[column_letter].width = adjusted_width

# Save the file
wb.save('/app/valores_test_import.xlsx')
print("Archivo de prueba creado: /app/valores_test_import.xlsx")