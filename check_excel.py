import openpyxl

# Check the Excel file content
wb = openpyxl.load_workbook('/app/valores_test_import.xlsx')
ws = wb.active

print("Sheet name:", ws.title)
print("Columns:")
for i, cell in enumerate(ws[1], 1):
    print(f"  {i}. {cell.value}")

print("\nData rows:")
for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
    print(f"Row {row_num}: {row}")